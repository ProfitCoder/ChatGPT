
# -*- coding: utf-8 -*-
"""
Conciliador Bancario Local - ULTRA 10 SEGURIDAD REFORZADAncipales:
- Marca SOLO la celda del importe: amarillo + subrayado.
- Archivos de salida: nombre_original_PUNTEADO.xlsx + INFORME_DUDAS_Y_PENDIENTES.xlsx.
- Puntea si:
  1) importe exacto,
  2) misma fecha,
  3) único candidato,
  4) y tercero válido compatible o concepto fuerte.
- También puntea por remesa Rel. Transferencia si la suma exacta coincide con único movimiento.
- Opción "Modo aprendizaje": permite puntear por importe+fecha únicos cuando no hay tercero claro.
- Ignora terceros genéricos: 99999998, 999999998, 99999999, deudores varios, etc.
- Evita cuelgues: agrupaciones por combinación desactivadas por defecto.
"""

import os, re, math, shutil, unicodedata, traceback, threading
from datetime import datetime
from itertools import combinations
from difflib import SequenceMatcher
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter

APP_NAME = "Conciliador Bancario Local - ULTRA 10 SEGURIDAD REFORZADA"

YELLOW = PatternFill("solid", fgColor="FFFF00")
HEADER = PatternFill("solid", fgColor="D9EAF7")
GREEN = PatternFill("solid", fgColor="E2F0D9")
ORANGE = PatternFill("solid", fgColor="F4B183")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

GENERIC_TERMS = {
    "99999998", "99999999", "999999998", "00000000", "SINNIF", "SINDNI", "SINNIE",
    "VARIOS", "DEUDORESVARIOS", "TERCEROSVARIOS", "ACREEDORESVARIOS",
    "DESCONOCIDO", "DESCONOCIDA", "NOIDENTIFICADO", "NOIDENTIFICADA"
}
STOPWORDS = {
    "DE","DEL","LA","LAS","EL","LOS","Y","A","EN","POR","PARA","CON",
    "PAGO","COBRO","TRANSFERENCIA","TRANSF","MOVIMIENTO","MOV","NUM","N",
    "AYTO","AYUNTAMIENTO","CUENTA","ORDINAL","FACTURA","FRA","RECIBO",
    "ABONO","CARGO","CONCEPTO","REFERENCIA","REF","BANCO","BANCARIA",
    "SEPA","OPERACION","OPERACIÓN","IMPORTE","FECHA","DIA","DÍA","SU",
    "SUS","AL","UN","UNA","S.L","SL","S.A","SA","S.L.","S.A.","MAS","MÁS",
    "DATOS","NOMBRE","TERCERO","LINEA","REL","RELACION","GENERICO","GENÉRICO","DUPLICADO","DUP"
}

def clean_text(x):
    if pd.isna(x): return ""
    text = str(x)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()
    text = re.sub(r"[^A-Z0-9/.\- ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text

def compact_text(x):
    return re.sub(r"[^A-Z0-9]+", "", clean_text(x))

def is_generic_tercero(x):
    c = compact_text(x)
    if not c: return True
    if c in GENERIC_TERMS: return True
    if c.startswith("99999998") or c.startswith("99999999"): return True
    if "DEUDORESVARIOS" in c or "TERCEROSVARIOS" in c or "ACREEDORESVARIOS" in c: return True
    return False

def to_float(x):
    if pd.isna(x): return None
    if isinstance(x, (int, float)):
        if isinstance(x, float) and math.isnan(x): return None
        return round(float(x), 2)
    txt = str(x).strip()
    if not txt: return None
    neg = False
    if txt.startswith("(") and txt.endswith(")"):
        neg = True; txt = txt[1:-1]
    txt = txt.replace("€","").replace("EUR","").replace(" ","").replace("+","")
    if txt.endswith("-"):
        neg = True; txt = txt[:-1]
    if "," in txt:
        txt = txt.replace(".","").replace(",",".")
    try:
        val = float(txt)
        return round(-val if neg else val, 2)
    except Exception:
        return None


def to_date(x):
    if pd.isna(x): return None
    if isinstance(x, datetime):
        d=x.date()
        return d if is_reasonable_date(d) else None
    if isinstance(x,(int,float)) and not isinstance(x,bool):
        try:
            if 20000 <= float(x) <= 60000:
                d=(datetime(1899,12,30)+pd.to_timedelta(float(x), unit="D")).date()
                return d if is_reasonable_date(d) else None
        except Exception: pass
    try:
        dt=pd.to_datetime(x, dayfirst=True, errors="coerce")
        if pd.isna(dt): return None
        d=dt.date()
        return d if is_reasonable_date(d) else None
    except Exception: return None

def is_reasonable_date(d):
    try:
        today=datetime.today().date()
        if d.year < 1990: return False
        if d > (today + pd.DateOffset(years=1)).date(): return False
        return True
    except Exception: return False

def same_day(a,b):
    return a is not None and b is not None and a == b

def amount_equal(a,b):
    return round(float(a)-float(b),2) == 0.0

def important_words(text):
    words = []
    for w in clean_text(text).split():
        w2 = w.strip(".-/")
        if len(w2) < 4: continue
        if w2 in STOPWORDS: continue
        if w2.isdigit():
            if len(w2) >= 5: words.append(w2)
            continue
        words.append(w2)
    return set(words)



REAL_REFERENCE_PREFIXES = ("FRA", "FAC", "FACT", "REF", "EXP", "OPA", "RC", "T", "REM", "REL", "NOM", "IRPF", "SS")

def is_real_reference_word(w):
    """
    ULTRA 10:
    No descartar referencias reales de facturas/expedientes.
    Ejemplos válidos: FRA12345, REF987654, EXP2024123, FAC-001, T/2026/300.
    """
    raw = str(w or "").strip().upper()
    compact = re.sub(r"[^A-Z0-9/]", "", raw)
    if len(compact) < 5:
        return False
    if any(compact.startswith(p) for p in REAL_REFERENCE_PREFIXES) and any(ch.isdigit() for ch in compact):
        return True
    if re.fullmatch(r"T/\d{4}/\d+", compact):
        return True
    return False

def is_generic_reference_word(w):
    w_raw = str(w or "").strip().upper()
    w = clean_text(w_raw).replace(" ", "")

    # Si parece referencia real, NO se considera genérica.
    if is_real_reference_word(w_raw) or is_real_reference_word(w):
        return False

    # Fechas, años sueltos, importes o códigos demasiado genéricos.
    if re.fullmatch(r"\d{4}", w):
        return True
    if re.fullmatch(r"\d{1,2}[/-]\d{1,2}[/-]\d{2,4}", w):
        return True
    if re.fullmatch(r"\d+[.,]\d+", w):
        return True
    if w in {"2026", "2025", "2024", "2023", "2022", "FEBRERO", "ENERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"}:
        return True
    return False


def concept_match_strong(a,b,min_common=2):
    wa, wb = important_words(a), important_words(b)
    if not wa or not wb: return False, 0, ""
    common = wa & wb
    common_clean = {w for w in common if not is_generic_reference_word(w)}
    has_good_long_ref = any(len(w) >= 8 and any(ch.isdigit() for ch in w) and not is_generic_reference_word(w) for w in common_clean)
    ok = len(common_clean) >= min_common or (has_good_long_ref and len(common_clean) >= 2)
    return ok, len(common_clean), ", ".join(sorted(common_clean))

def text_score(a,b):
    wa, wb = important_words(a), important_words(b)
    if not wa or not wb: return 0.0
    return len(wa & wb) / max(1, len(wa | wb))

def extract_refs(text):
    t = clean_text(text)
    refs = set()
    patterns = [
        r"T/\d{4}/\d+", r"SW-\d{4}-\d+", r"B-\d+",
        r"FRA\.?\s*[A-Z0-9/\-]+", r"FACTURA\s*[A-Z0-9/\-]+",
        r"RECIBO\s*[A-Z0-9/\-]+", r"EXP\.?\s*[A-Z0-9/\-]+",
        r"[XYZ]\d{7}[A-Z]", r"[ABCDEFGHJKLMNPQRSUVW]\d{7}[A-Z0-9]",
        r"\d{8}[A-Z]"
    ]
    for p in patterns:
        for m in re.findall(p, t): refs.add(m.strip())
    return sorted(refs)


TERCERO_STOPWORDS = {
    "SL","SA","SLL","SLU","SAU","SLC","SOCIEDAD","LIMITADA","ANONIMA","ANÓNIMA",
    "EMPRESA","SERVICIOS","GRUPO","THE","DE","DEL","LA","LAS","EL","LOS","Y","&",
    "S.L","S.A","S.L.","S.A.","U","ESPANA","ESPAÑA"
}

def tercero_tokens(x):
    toks = []
    for w in clean_text(x).replace('.', ' ').replace('-', ' ').split():
        w = w.strip()
        if len(w) < 4: continue
        if w in TERCERO_STOPWORDS: continue
        if w.isdigit(): continue
        toks.append(w)
    return set(toks)


def tercero_tokens(x):
    toks = set(important_words(x))
    # eliminar palabras jurídicas/comunes que no distinguen bien
    generic = {"SOCIEDAD", "LIMITADA", "ANONIMA", "ANÓNIMA", "EMPRESA", "SERVICIOS", "GRUPO", "SL", "SA", "SLL", "SLA", "U", "UTE"}
    return {t for t in toks if t not in generic and len(t) >= 4}

def same_tercero(a,b):
    """
    ULTRA 10:
    Más estricto. Evita falsos positivos por una sola palabra genérica compartida.
    """
    ca, cb = clean_text(a), clean_text(b)
    if not ca or not cb:
        return False
    if is_generic_tercero(ca) or is_generic_tercero(cb):
        return False

    if ca == cb:
        return True

    ta, tb = tercero_tokens(ca), tercero_tokens(cb)
    if not ta or not tb:
        return False

    common = ta & tb

    # Dos palabras distintivas comunes = fuerte.
    if len(common) >= 2:
        return True

    # Una sola palabra solo vale si es muy característica y ambos textos no son demasiado genéricos.
    if len(common) == 1:
        w = next(iter(common))
        if len(w) >= 8:
            # Evitar que "SERVICIOS", "MUNICIPAL", etc. decidan solos.
            banned_single = {"SERVICIOS", "MUNICIPAL", "MUNICIPALES", "PROVEEDOR", "CLIENTES", "ENERGIA", "ENERGÍA"}
            if w not in banned_single:
                return True

    # Subcadena solo si el texto corto tiene al menos dos tokens distintivos.
    shorter, longer = (ca, cb) if len(ca) < len(cb) else (cb, ca)
    stoks = tercero_tokens(shorter)
    if len(stoks) >= 2 and shorter in longer:
        return True

    return False


def detect_header_row(path, sheet_name=0):
    raw = pd.read_excel(path, sheet_name=sheet_name, header=None, dtype=object)
    keywords = ["AYUNTAMIENTO","CUENTA","ORDINAL","FECHA","CONCILIACION","CONCILIACIÓN",
                "OPERACION","OPERACIÓN","CONCEPTO","IMPORTE","SALDO","REFERENCIA",
                "DEBE","HABER","INGRESOS","PAGOS","TERCERO","PROVEEDOR","NOMBRE",
                "RAZON","RAZÓN","NIF","N.I.F","DNI","NIE","CIF","MAS DATOS","MÁS DATOS"]
    best_row, best_score = 0, -1
    for i in range(min(60, len(raw))):
        row_txt = " ".join(clean_text(v) for v in raw.iloc[i].tolist() if pd.notna(v))
        score = sum(1 for k in keywords if clean_text(k) in row_txt)
        score += min(raw.iloc[i].notna().sum(), 15) * 0.1
        if score > best_score:
            best_score, best_row = score, i
    return best_row

def read_excel(path, sheet_name=0):
    header = detect_header_row(path, sheet_name=sheet_name)
    df = pd.read_excel(path, sheet_name=sheet_name, header=header, dtype=object)
    df = df.dropna(how="all").copy()
    df.columns = [str(c).strip() for c in df.columns]
    return df, header

def find_col(df, names):
    cols = {c: clean_text(c) for c in df.columns}
    for name in names:
        n = clean_text(name)
        for c, cc in cols.items():
            if n == cc: return c
    for name in names:
        n = clean_text(name)
        for c, cc in cols.items():
            if n and n in cc: return c
    return None

def find_cols(df, names):
    cols = []
    cleaned = {c: clean_text(c) for c in df.columns}
    for name in names:
        n = clean_text(name)
        for c, cc in cleaned.items():
            if n and (n == cc or n in cc) and c not in cols:
                cols.append(c)
    return cols



PARTIAL_AMOUNT_BAD_WORDS = {
    "IVA", "RETENCION", "RETENCIÓN", "IRPF", "COMISION", "COMISIÓN", "BASE", "IMPONIBLE",
    "DESCUENTO", "DTO", "RECARGO", "INTERES", "INTERÉS", "TASA", "PORCENTAJE", "%",
    "CUOTA", "CUOTA IVA", "GASTO", "GASTOS", "BONIFICACION", "BONIFICACIÓN"
}
SPLIT_INGRESO_NAMES = ["Ingresos", "Ingreso", "Cobros", "Cobro", "Abonos", "Abono", "Debe"]
SPLIT_PAGO_NAMES = ["Pagos", "Pago", "Cargos", "Cargo", "Haber"]

def is_partial_amount_col(col_name):
    c = clean_text(col_name)
    return any(bad in c for bad in PARTIAL_AMOUNT_BAD_WORDS)

def amount_col_score(col_name):
    c = clean_text(col_name)
    exact_good = {
        "IMPORTE": 100, "IMPORTE MOVIMIENTO": 98, "IMPORTE CONCILIACION": 98,
        "IMPORTE CONCILIACIÓN": 98, "IMPORTE EUR": 96, "TOTAL": 90,
        "LIQUIDO": 85, "LÍQUIDO": 85
    }
    if c in exact_good: return exact_good[c], False
    if is_partial_amount_col(c): return 10, True
    if "IMPORTE" in c: return 55, False
    if "TOTAL" in c: return 50, False
    return 0, False

def find_amount_candidates(df):
    candidates=[]
    for c in df.columns:
        score,partial=amount_col_score(c)
        numeric_count=sum(to_float(v) is not None for v in df[c].head(500))
        if score > 0 and numeric_count:
            score += min(numeric_count,50)*0.1
        if score>0 and numeric_count>0: candidates.append((c,score,partial,numeric_count))
    candidates.sort(key=lambda x:x[1], reverse=True)
    return candidates

def split_col_candidates(df, names):
    """Candidatas a Ingresos/Pagos filtrando columnas parciales."""
    cols=[]
    cleaned={c: clean_text(c) for c in df.columns}
    for name in names:
        n=clean_text(name)
        for c,cc in cleaned.items():
            if not n: continue
            if n == cc or n in cc:
                numeric_count=sum(to_float(v) is not None for v in df[c].head(500))
                if numeric_count <= 0: continue
                partial=is_partial_amount_col(cc)
                # Exacto preferido. Contiene pero parcial = sospechoso.
                score=(100 if n == cc else 60) + min(numeric_count,50)*0.1
                if partial: score=5
                cols.append((c, score, partial, numeric_count, name))
    # Unificar por columna con mejor score.
    best={}
    for item in cols:
        c=item[0]
        if c not in best or item[1] > best[c][1]: best[c]=item
    res=list(best.values())
    res.sort(key=lambda x:x[1], reverse=True)
    return res

def choose_split_col(df, names):
    cands=split_col_candidates(df, names)
    safe=[x for x in cands if not x[2]]
    if safe:
        return safe[0][0], cands
    return None, cands

def find_amount_mode(df):
    ingresos, ingresos_cands = choose_split_col(df, SPLIT_INGRESO_NAMES)
    pagos, pagos_cands = choose_split_col(df, SPLIT_PAGO_NAMES)
    if ingresos or pagos:
        return "split", ingresos, pagos
    candidates=find_amount_candidates(df)
    if candidates: return "importe", candidates[0][0], None
    numeric_cols=[]
    for c in df.columns:
        count=sum(to_float(v) is not None for v in df[c].head(500))
        if count>0: numeric_cols.append((c,count))
    numeric_cols.sort(key=lambda x:x[1], reverse=True)
    return ("importe", numeric_cols[0][0], None) if numeric_cols else ("importe", None, None)

def normalize(path, name, sheet_name=0):
    df, header = read_excel(path, sheet_name=sheet_name)
    col_pos = {c: i+1 for i, c in enumerate(df.columns)}
    ayto_col = find_col(df, ["Ayuntamiento","Ayto"])
    ordinal_col = find_col(df, ["Cuenta/Ordinal","Ordinal","Cuenta"])
    fecha_col = find_col(df, ["Fecha Conciliación","Fecha Conciliacion","Fecha operación","Fecha Operacion","Fecha valor","Fecha"])
    concepto_cols = find_cols(df, ["Concepto banco","Texto Explicativo","Texto","Concepto","Movimiento","Descripcion","Descripción","Más datos","Mas datos"])
    tercero_col = find_col(df, ["Nombre Tercero","Nombre tercero","Razón social","Razon social","Proveedor","Beneficiario","Interesado","Tercero","N.I.F. Tercero","NIF","DNI","NIE","CIF"])
    nif_col = find_col(df, ["N.I.F. Tercero","NIF","DNI","NIE","CIF"])
    referencia_col = find_col(df, ["Referencia banco","Referencia","Nº operación","N° operación","Numero operacion","Rel. Transferencia"])
    rel_col = find_col(df, ["Rel. Transferencia","Rel Transferencia"])
    amount_mode, amount_a, amount_b = find_amount_mode(df)

    rows = []
    for idx, row in df.iterrows():
        fecha = to_date(row.get(fecha_col)) if fecha_col else None
        importe, amount_col_num = None, None
        if amount_mode == "split":
            ingreso = to_float(row.get(amount_a)) if amount_a else None
            pago = to_float(row.get(amount_b)) if amount_b else None
            if ingreso is not None and abs(ingreso) > 0:
                importe, amount_col_num = abs(ingreso), col_pos.get(amount_a)
            elif pago is not None and abs(pago) > 0:
                importe, amount_col_num = -abs(pago), col_pos.get(amount_b)
        else:
            importe = to_float(row.get(amount_a)) if amount_a else None
            amount_col_num = col_pos.get(amount_a)
        if importe is None or amount_col_num is None: continue

        ayto = clean_text(row.get(ayto_col)) if ayto_col else ""
        ordinal = clean_text(row.get(ordinal_col)) if ordinal_col else ""
        tercero = clean_text(row.get(tercero_col)) if tercero_col else ""
        nif = clean_text(row.get(nif_col)) if nif_col else ""
        parts = []
        for c in concepto_cols + [tercero_col, nif_col, referencia_col]:
            if c:
                v = row.get(c)
                if pd.notna(v): parts.append(str(v))
        if not parts:
            for c in df.columns:
                if c in [fecha_col, amount_a, amount_b]: continue
                v = row.get(c)
                if pd.notna(v): parts.append(str(v))
        concepto = " ".join(parts).strip()
        concepto_limpio = clean_text(concepto)
        refs = extract_refs(concepto_limpio)
        rel = ""
        if rel_col and pd.notna(row.get(rel_col)): rel = clean_text(row.get(rel_col))
        if not rel:
            for ref in refs:
                if re.match(r"T/\d{4}/\d+", ref):
                    rel = ref
                    break
        rows.append({
            "name": name, "excel_row": int(header + idx + 2), "amount_col": int(amount_col_num),
            "ayuntamiento": ayto, "ordinal": ordinal, "fecha": fecha, "importe": round(float(importe),2),
            "tercero": tercero, "nif": nif, "tercero_generico": is_generic_tercero(tercero) or is_generic_tercero(nif),
            "concepto": concepto, "concepto_limpio": concepto_limpio,
            "palabras_importantes": ", ".join(sorted(important_words(concepto_limpio))),
            "refs": ", ".join(refs), "relacion": rel,
        })
    return pd.DataFrame(rows), {
        "header_row": header+1, "ayto_col": ayto_col, "ordinal_col": ordinal_col, "fecha_col": fecha_col,
        "tercero_col": tercero_col, "nif_col": nif_col, "amount_mode": amount_mode,
        "amount_a": amount_a, "amount_b": amount_b, "concepto_cols": ", ".join(concepto_cols), "rel_col": rel_col
    }


def scope_status(a,b):
    """
    ULTRA 10: si un ámbito existe en un lado y falta en el otro, no se da por compatible.
    Si no hay ningún dato de ámbito en ninguno de los dos lados, se permite, pero queda no verificable.
    """
    ay_a, ay_b = a.get("ayuntamiento", ""), b.get("ayuntamiento", "")
    or_a, or_b = a.get("ordinal", ""), b.get("ordinal", "")
    has_any = bool(ay_a or ay_b or or_a or or_b)
    if not has_any:
        return True, "ámbito no informado en ambos lados"
    if ay_a and ay_b and ay_a != ay_b:
        return False, "Ayuntamiento distinto"
    if or_a and or_b and or_a != or_b:
        return False, "Ordinal distinto"
    if (ay_a and not ay_b) or (ay_b and not ay_a):
        return False, "Ayuntamiento no verificable"
    if (or_a and not or_b) or (or_b and not or_a):
        return False, "Ordinal no verificable"
    return True, "ámbito correcto"

def same_scope(a,b):
    return scope_status(a,b)[0]



def review_confidence_and_risk(tipo):
    t=clean_text(tipo)
    if "VARIOS" in t or "COMPARTIDO" in t or "AMBIGUA" in t: return 40,"Alto"
    if "DIFERENCIA" in t or "FECHA DISTINTA" in t: return 55,"Medio"
    if "POSIBLE AGRUPACION" in t or "AGRUPACION" in t: return 65,"Medio"
    if "PENDIENTE" in t: return 0,"Pendiente"
    return 50,"Medio"

def add_review(rows, tipo, left_rows, right_rows, motivo, que_revisar):
    conf,risk=review_confidence_and_risk(tipo)
    rows.append({"Estado":tipo,"Confianza":conf,"Riesgo":risk,"Motivo resumido":motivo[:140],"Filas Excel 1":left_rows,"Filas Excel 2":right_rows,"Motivo":motivo,"Qué revisar":que_revisar})



def audit_candidate(li, ri, left, right, left_candidates, right_candidates, motivo_base, strong_unique=False):
    lr, rr = left.loc[li], right.loc[ri]
    checks=[]; risk=0
    if not amount_equal(lr["importe"], rr["importe"]): checks.append("importe no exacto"); risk += 100
    if not same_day(lr["fecha"], rr["fecha"]): checks.append("fecha distinta"); risk += 100
    scope_ok, scope_msg = scope_status(lr, rr)
    if not scope_ok: checks.append(scope_msg); risk += 100
    elif scope_msg == "ámbito no informado en ambos lados":
        checks.append(scope_msg)
        # Aviso leve: no bloquea si el resto es muy claro.
        risk += 5

    lc=[x for x in left_candidates.get(li,[]) if not right.loc[x,"matched"]]
    rc=[x for x in right_candidates.get(ri,[]) if not left.loc[x,"matched"]]

    # ULTRA 10: si ya hay tercero/concepto fuerte y único en ambos sentidos, no castigamos por
    # otros movimientos con mismo importe/fecha que no son coincidencias fuertes.
    if not strong_unique:
        if len(lc)!=1: checks.append(f"Excel1 tiene {len(lc)} candidatos"); risk += 40
        if len(rc)!=1: checks.append(f"Excel2 tiene {len(rc)} candidatos"); risk += 40
    else:
        if len(lc)!=1 or len(rc)!=1:
            checks.append("hay duplicados de importe/fecha, pero el tercero/concepto fuerte es único")
            risk += 5

    if lr.get("tercero_generico") and rr.get("tercero_generico"):
        checks.append("terceros genéricos")
        risk += 5
    if risk==0: return True,100,"Bajo",motivo_base+" | auditoría correcta"
    if risk<=15: return True,90,"Bajo",motivo_base+" | auditoría con aviso leve: "+", ".join(checks)
    return False,max(0,100-risk),"Alto","No marcado por auditoría: "+", ".join(checks)

def reconcile(left, right, allow_relation_groups=True, allow_unique_date_amount=True, allow_combo_groups=False, min_common_words=2):
    left,right=left.copy(),right.copy(); left["matched"]=False; right["matched"]=False
    secure=[]; review=[]; reviewed_left=set(); reviewed_right=set()
    def add_secure(tipo, confianza, lis, ris, motivo):
        total_l=round(float(left.loc[lis,"importe"].sum()),2); total_r=round(float(right.loc[ris,"importe"].sum()),2)
        left.loc[lis,"matched"]=True; right.loc[ris,"matched"]=True
        secure.append({"tipo":tipo,"confianza":confianza,"left_indices":list(lis),"right_indices":list(ris),"importe_excel1":total_l,"importe_excel2":total_r,"diferencia":round(total_l-total_r,2),"motivo":motivo})
    def add_review_once(tipo,left_ids,right_ids,motivo,que_revisar):
        for i in left_ids: reviewed_left.add(i)
        for i in right_ids: reviewed_right.add(i)
        add_review(review,tipo,", ".join(str(int(left.loc[i,"excel_row"])) for i in left_ids),", ".join(str(int(right.loc[i,"excel_row"])) for i in right_ids),motivo,que_revisar)
    def compatible_amount_date(li,ri):
        lr,rr=left.loc[li],right.loc[ri]
        return same_scope(lr,rr) and amount_equal(lr["importe"],rr["importe"]) and same_day(lr["fecha"],rr["fecha"])
    def strong_match(li,ri):
        lr,rr=left.loc[li],right.loc[ri]
        tercero_ok=same_tercero(lr.get("tercero",""),rr.get("tercero","")) or same_tercero(lr.get("tercero",""),rr.get("concepto","")) or same_tercero(rr.get("tercero",""),lr.get("concepto",""))
        concepto_ok,_,common_words=concept_match_strong(lr.get("concepto_limpio",""),rr.get("concepto_limpio",""),min_common=min_common_words)
        if tercero_ok: return True,"Punteado seguro por tercero",100,"tercero compatible"
        if concepto_ok: return True,"Punteado seguro por concepto",92,f"concepto fuerte: {common_words}"
        return False,"",0,""
    # Remesas: suma + fecha + mismo ámbito
    if allow_relation_groups:
        for side in ("left","right"):
            source=left if side=="left" else right; target=right if side=="left" else left
            rel_groups={}
            for si,sr in source[~source["matched"]].iterrows():
                rel=str(sr.get("relacion","")).strip()
                if rel and rel.upper() not in ("NONE","NAN",""): rel_groups.setdefault(rel,[]).append(si)
            for rel,sis in rel_groups.items():
                total=round(float(source.loc[sis,"importe"].sum()),2); fechas=sorted({x for x in source.loc[sis,"fecha"].tolist() if x is not None}); cands=[]
                for ti,tr in target[~target["matched"]].iterrows():
                    if not amount_equal(total,tr["importe"]): continue
                    if fechas and tr["fecha"] not in fechas: continue
                    scope_ok=True
                    for si in sis:
                        if not same_scope(source.loc[si], tr): scope_ok=False; break
                    if scope_ok: cands.append(ti)
                if len(cands)==1:
                    if side=="left": add_secure("Punteado seguro por remesa auditada",98,sis,[cands[0]],f"Relación {rel}: suma exacta, fecha compatible, ámbito correcto y único candidato")
                    else: add_secure("Punteado seguro por remesa auditada",98,[cands[0]],sis,f"Relación {rel}: suma exacta, fecha compatible, ámbito correcto y único candidato")
                elif len(cands)>1:
                    if side=="left": add_review_once("DUDA - remesa con varios candidatos",sis,cands[:10],f"Relación {rel} suma {total}, pero hay varios candidatos del mismo ámbito.","Elegir manualmente.")
                    else: add_review_once("DUDA - remesa con varios candidatos",cands[:10],sis,f"Relación {rel} suma {total}, pero hay varios candidatos del mismo ámbito.","Elegir manualmente.")
    # candidates index
    right_index={}
    for ri,rr in right[~right["matched"]].iterrows():
        amount=round(float(rr.get("importe")),2); right_index.setdefault((rr.get("fecha"),amount,rr.get("ayuntamiento",""),rr.get("ordinal","")),[]).append(ri); right_index.setdefault((rr.get("fecha"),amount,"",""),[]).append(ri)
    left_candidates={}; right_candidates={}; strong_left_candidates={}; strong_right_candidates={}
    for li,lr in left[~left["matched"]].iterrows():
        amount=round(float(lr.get("importe")),2); seen=set(); possible=[]
        for key in [(lr.get("fecha"),amount,lr.get("ayuntamiento",""),lr.get("ordinal","")),(lr.get("fecha"),amount,"","")]:
            for ri in right_index.get(key,[]):
                if ri in seen: continue
                seen.add(ri)
                if not right.loc[ri,"matched"] and compatible_amount_date(li,ri): possible.append(ri)
        for ri in possible:
            left_candidates.setdefault(li,[]).append(ri); right_candidates.setdefault(ri,[]).append(li)
            ok,tipo,conf,motivo=strong_match(li,ri)
            if ok: strong_left_candidates.setdefault(li,[]).append((ri,tipo,conf,motivo)); strong_right_candidates.setdefault(ri,[]).append((li,tipo,conf,motivo))
    # strong audited
    for li in list(left[~left["matched"]].index):
        cands=[x for x in strong_left_candidates.get(li,[]) if not right.loc[x[0],"matched"] and not left.loc[li,"matched"]]
        if len(cands)!=1: continue
        ri,tipo,conf,motivo=cands[0]
        reverse=[x for x in strong_right_candidates.get(ri,[]) if not left.loc[x[0],"matched"] and not right.loc[ri,"matched"]]
        if len(reverse)==1 and reverse[0][0]==li:
            ok_audit,audit_conf,risk,audit_motive=audit_candidate(li,ri,left,right,left_candidates,right_candidates,f"Importe exacto, misma fecha y {motivo}", strong_unique=True)
            if ok_audit: add_secure(tipo+" auditado",min(conf,audit_conf),[li],[ri],audit_motive)
            else: add_review_once("DUDA - auditoría rechazó punteo",[li],[ri],audit_motive,"Revisar manualmente.")
    # amount/date unique audited
    if allow_unique_date_amount:
        for li in list(left[~left["matched"]].index):
            cands=[ri for ri in left_candidates.get(li,[]) if not right.loc[ri,"matched"] and not left.loc[li,"matched"]]
            if len(cands)!=1: continue
            ri=cands[0]
            reverse=[lj for lj in right_candidates.get(ri,[]) if not left.loc[lj,"matched"] and not right.loc[ri,"matched"]]
            if len(reverse)==1 and reverse[0]==li:
                ok_audit,audit_conf,risk,audit_motive=audit_candidate(li,ri,left,right,left_candidates,right_candidates,"Importe exacto y misma fecha con candidato único bidireccional", strong_unique=False)
                if ok_audit: add_secure("Punteado seguro por importe y fecha únicos auditado",min(90,audit_conf),[li],[ri],audit_motive)
                else: add_review_once("DUDA - auditoría rechazó punteo",[li],[ri],audit_motive,"Revisar manualmente.")
    # ambiguity
    for li in left[~left["matched"]].index:
        cands=[ri for ri in left_candidates.get(li,[]) if not right.loc[ri,"matched"]]
        if len(cands)>1: add_review_once("DUDA - varios candidatos por importe/fecha",[li],cands[:10],"Mismo importe y fecha con varios candidatos.","Elegir manualmente.")
        elif len(cands)==1:
            ri=cands[0]; reverse=[lj for lj in right_candidates.get(ri,[]) if not left.loc[lj,"matched"]]
            if len(reverse)>1: add_review_once("DUDA - candidato compartido",reverse[:10],[ri],"Una misma línea del Excel 2 puede corresponder a varias líneas del Excel 1.","No se marca porque no es único en ambos sentidos.")
    # date near indexed
    right_by_amount={}
    for ri,rr in right[~right["matched"]].iterrows(): right_by_amount.setdefault(round(float(rr["importe"]),2),[]).append(ri)
    for li,lr in left[~left["matched"]].iterrows():
        if li in reviewed_left: continue
        cands=[]
        for ri in right_by_amount.get(round(float(lr["importe"]),2),[]):
            rr=right.loc[ri]
            if not same_scope(lr,rr) or same_day(lr["fecha"],rr["fecha"]): continue
            if lr["fecha"] is not None and rr["fecha"] is not None:
                dd=abs((lr["fecha"]-rr["fecha"]).days)
                if dd<=10: cands.append((ri,dd))
        if len(cands)==1: add_review_once("DUDA - fecha distinta",[li],[cands[0][0]],f"Mismo importe, pero fecha distinta por {cands[0][1]} día(s).","No puntear automático.")
        elif len(cands)>1: add_review_once("DUDA - mismo importe en fechas cercanas",[li],[i for i,_ in cands[:10]],"Mismo importe en fechas cercanas.","Revisar.")
    # cent diff indexed by date
    right_by_date={}
    for ri,rr in right[~right["matched"]].iterrows(): right_by_date.setdefault(rr["fecha"],[]).append(ri)
    for li,lr in left[~left["matched"]].iterrows():
        if li in reviewed_left: continue
        cands=[]
        for ri in right_by_date.get(lr["fecha"],[]):
            rr=right.loc[ri]
            if not same_scope(lr,rr): continue
            diff=abs(round(float(lr["importe"])-float(rr["importe"]),2))
            if diff==0 or diff>1.0: continue
            concepto_ok,_,_=concept_match_strong(lr["concepto_limpio"],rr["concepto_limpio"],min_common=min_common_words)
            if concepto_ok: cands.append((ri,diff))
        if len(cands)==1: add_review_once("DUDA - diferencia de importe",[li],[cands[0][0]],f"Diferencia de {cands[0][1]} €. Todo debe cuadrar al céntimo.","Corregir o revisar.")
    # combo optional
    if allow_combo_groups:
        for li,lr in left[~left["matched"]].iterrows():
            if li in reviewed_left: continue
            pool=[ri for ri,rr in right[~right["matched"]].iterrows() if same_scope(lr,rr) and lr["importe"]*rr["importe"]>0 and same_day(lr["fecha"],rr["fecha"])]
            if 2<=len(pool)<=12:
                found=False
                for r in range(2,min(4,len(pool))+1):
                    for combo in combinations(pool,r):
                        if amount_equal(round(float(right.loc[list(combo),"importe"].sum()),2),lr["importe"]): add_review_once("DUDA - posible agrupación",[li],list(combo),f"{r} importes del Excel 2 suman exactamente el importe del Excel 1.","Confirmar remesa."); found=True; break
                    if found: break
    for li,lr in left[~left["matched"]].iterrows():
        if li not in reviewed_left: add_review_once("PENDIENTE Excel 1",[li],[],f"No tiene pareja segura. Importe: {lr['importe']} Fecha: {lr['fecha']} Tercero: {lr.get('tercero','')}","Revisar.")
    for ri,rr in right[~right["matched"]].iterrows():
        if ri not in reviewed_right: add_review_once("PENDIENTE Excel 2",[],[ri],f"No tiene pareja segura. Importe: {rr['importe']} Fecha: {rr['fecha']} Tercero: {rr.get('tercero','')}","Revisar.")
    return secure,review,left,right

def output_name(path):
    stem, ext = os.path.splitext(os.path.basename(path))
    if stem.upper().endswith("_PUNTEADO"):
        stem = stem[:-9]
    return f"{stem}_PUNTEADO.xlsx"

def safe_output_path(out_dir, input_path):
    base_name = output_name(input_path)
    candidate = os.path.join(out_dir, base_name)
    if not os.path.exists(candidate):
        return candidate
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem, ext = os.path.splitext(base_name)
    return os.path.join(out_dir, f"{stem}_{stamp}{ext}")

def safe_named_output_path(out_dir, filename):
    candidate=os.path.join(out_dir, filename)
    if not os.path.exists(candidate): return candidate
    stamp=datetime.now().strftime("%Y%m%d_%H%M%S")
    stem,ext=os.path.splitext(filename)
    return os.path.join(out_dir, f"{stem}_{stamp}{ext}")


def mark_only_amounts(input_path, output_path, marks):
    ext = os.path.splitext(input_path)[1].lower()
    if ext in (".xlsx",".xlsm"):
        shutil.copy2(input_path, output_path)
        wb = load_workbook(output_path)
        ws = wb[wb.sheetnames[0]]
    else:
        df, header = read_excel(input_path)
        wb = Workbook(); ws = wb.active; ws.title = "PUNTEADO"
        for c, col in enumerate(df.columns,1):
            ws.cell(1,c,str(col)).fill = HEADER; ws.cell(1,c).font = Font(bold=True)
        for r, (_, row) in enumerate(df.iterrows(),2):
            for c, col in enumerate(df.columns,1):
                ws.cell(r,c,row[col])
        marks = [(excel_row - header, amount_col) for excel_row, amount_col in marks]
    for excel_row, amount_col in marks:
        if 1 <= excel_row <= ws.max_row and 1 <= amount_col <= ws.max_column:
            cell = ws.cell(row=excel_row, column=amount_col)
            cell.fill = YELLOW
            old = cell.font
            cell.font = Font(name=old.name, sz=old.sz, bold=old.bold, italic=old.italic,
                             underline="single", strike=old.strike, color=old.color)
    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            max_len = max(max_len, len("" if cell.value is None else str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len+2,10),65)
    wb.save(output_path)

def autofit(ws):
    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            max_len = max(max_len, len("" if cell.value is None else str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len+2,10),90)

def write_df(ws, df, body_fill=None):
    if df is None or df.empty:
        ws.cell(1,1,"Sin datos"); return
    for c, col in enumerate(df.columns,1):
        cell=ws.cell(1,c,col); cell.fill=HEADER; cell.font=Font(bold=True)
    for r, (_, row) in enumerate(df.iterrows(),2):
        for c, col in enumerate(df.columns,1):
            val=row[col]
            if hasattr(val, "isoformat"): val=val.isoformat()
            cell=ws.cell(r,c,val)
            if body_fill: cell.fill=body_fill
    for row in ws.iter_rows():
        for cell in row:
            cell.border=BORDER; cell.alignment=Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes="A2"; autofit(ws)

def make_report(path, secure, review, left, right, name_left, name_right, meta_left, meta_right, min_common_words):
    wb=Workbook(); ws=wb.active; ws.title="Resumen"
    resumen=pd.DataFrame([
        ["Excel 1", name_left], ["Excel 2", name_right],
        ["Movimientos Excel 1 leídos", len(left)], ["Movimientos Excel 2 leídos", len(right)],
        ["Punteados seguros", len(secure)], ["Registros en Dudas y pendientes", len(review)],
        ["Regla por tercero/concepto", f"Importe exacto + misma fecha + único candidato + tercero compatible o {min_common_words} palabras importantes"],
        ["Regla por importe/fecha únicos", "Activable: puntea si es único aunque no haya tercero claro"],
        ["Regla por remesa", "Rel. Transferencia: suma exacta contra único movimiento"],
        ["Terceros ignorados", "99999998, 999999998, deudores varios, etc."],
        ["Regla visual", "Solo celda de importe en amarillo + subrayado"],
        ["Columnas Excel 1", f"Fecha={meta_left.get('fecha_col')} | Tercero={meta_left.get('tercero_col')} | Importe={meta_left.get('amount_a') or meta_left.get('amount_b')} | Concepto={meta_left.get('concepto_cols')}"],
        ["Columnas Excel 2", f"Fecha={meta_right.get('fecha_col')} | Tercero={meta_right.get('tercero_col')} | Importe={meta_right.get('amount_a') or meta_right.get('amount_b')} | Concepto={meta_right.get('concepto_cols')}"],
    ], columns=["Concepto","Valor"])
    write_df(ws,resumen)
    rows=[]
    for s in secure:
        rows.append({"Estado":s["tipo"],"Confianza":s["confianza"],f"Importe {name_left}":s["importe_excel1"],f"Importe {name_right}":s["importe_excel2"],
                     "Diferencia":s["diferencia"],f"Filas {name_left}":", ".join(str(int(left.loc[i,"excel_row"])) for i in s["left_indices"]),
                     f"Filas {name_right}":", ".join(str(int(right.loc[i,"excel_row"])) for i in s["right_indices"]),"Motivo":s["motivo"]})
    ws2=wb.create_sheet("Punteados seguros"); write_df(ws2,pd.DataFrame(rows),GREEN)
    ws3=wb.create_sheet("Dudas y pendientes"); write_df(ws3,pd.DataFrame(review),ORANGE)
    manual=pd.DataFrame([
        ["1","Usa banco completo y SicalWin completo."],
        ["2","Abre primero Resumen para ver columnas detectadas."],
        ["3","Los *_PUNTEADO.xlsx marcan solo importes seguros."],
        ["4","Amarillo + subrayado = punteado seguro."],
        ["5","Si hay varios candidatos o céntimos de diferencia, va al informe."],
        ["6","Deja activadas remesas. Combinaciones solo si hace falta porque puede tardar."],
    ], columns=["Paso","Explicación"])
    ws4=wb.create_sheet("Manual"); write_df(ws4,manual)
    wb.save(path)

def process(file_left,file_right,out_dir,name_left,name_right,allow_relation_groups,allow_unique_date_amount,allow_combo_groups,min_common_words,sheet_left=0,sheet_right=0):
    out_dir = ensure_safe_output_dir(out_dir)
    os.makedirs(out_dir, exist_ok=True)

    backup_left = make_backup_copy(file_left, out_dir)
    backup_right = make_backup_copy(file_right, out_dir)

    left, meta_left = normalize(file_left,name_left,sheet_left)
    right, meta_right = normalize(file_right,name_right,sheet_right)
    secure, review, left2, right2 = reconcile(left,right,allow_relation_groups,allow_unique_date_amount,allow_combo_groups,min_common_words)
    left_marks=[]; right_marks=[]
    for s in secure:
        for i in s["left_indices"]: left_marks.append((int(left2.loc[i,"excel_row"]), int(left2.loc[i,"amount_col"])))
        for i in s["right_indices"]: right_marks.append((int(right2.loc[i,"excel_row"]), int(right2.loc[i,"amount_col"])))
    left_out=safe_output_path(out_dir,file_left)
    right_out=safe_output_path(out_dir,file_right)
    report=safe_named_output_path(out_dir,"INFORME_DUDAS_Y_PENDIENTES.xlsx")
    diagnostic=os.path.join(out_dir,"DIAGNOSTICO_CONCILIADOR_ULTRA10.txt")
    write_diagnostic_report(diagnostic, {
        "Excel 1": file_left,
        "Excel 2": file_right,
        "Copia seguridad Excel 1": backup_left,
        "Copia seguridad Excel 2": backup_right,
        "Movimientos Excel 1": len(left2),
        "Movimientos Excel 2": len(right2),
        "Punteados seguros": len(secure),
        "Registros informe": len(review),
        "Columnas Excel 1": meta_left,
        "Columnas Excel 2": meta_right,
        "Reglas": "ULTRA 10 IA reforzada: split protegido, tercero estricto, scope verificable, auditoría flexible para coincidencias fuertes",
    })
    mark_only_amounts(file_left,left_out,left_marks)
    mark_only_amounts(file_right,right_out,right_marks)
    make_report(report,secure,review,left2,right2,name_left,name_right,meta_left,meta_right,min_common_words)
    return {"left_out":left_out,"right_out":right_out,"report":report,"diagnostic":diagnostic,"left_count":len(left2),"right_count":len(right2),"secure":len(secure),"review":len(review),"meta_left":meta_left,"meta_right":meta_right}


def list_sheets(path):
    try:
        xls = pd.ExcelFile(path)
        return xls.sheet_names
    except Exception:
        return []

def validate_before_process(file_left, file_right, sheet_left=0, sheet_right=0):
    left, meta_left = normalize(file_left, "Excel 1", sheet_left)
    right, meta_right = normalize(file_right, "Excel 2", sheet_right)
    problems = []
    warnings = []

    for side_name, df, meta in [("Excel 1", left, meta_left), ("Excel 2", right, meta_right)]:
        if df.empty:
            problems.append(f"{side_name}: no se han leído movimientos con importe.")
        if not meta.get("fecha_col"):
            problems.append(f"{side_name}: no se ha detectado columna de fecha.")
        if not (meta.get("amount_a") or meta.get("amount_b")):
            problems.append(f"{side_name}: no se ha detectado columna de importe.")
        if not meta.get("concepto_cols") and not meta.get("tercero_col"):
            warnings.append(f"{side_name}: no se ha detectado concepto ni tercero; la conciliación será muy limitada.")
        suspicious, col = detect_suspicious_amount_column(meta)
        if suspicious:
            problems.append(f"{side_name}: la columna de importe detectada parece sospechosa: {col}. Selecciona/revisa el Excel antes de procesar.")
        try:
            raw_df, _ = read_excel(file_left if side_name == "Excel 1" else file_right, sheet_left if side_name == "Excel 1" else sheet_right)
            candidates = find_amount_candidates(raw_df)
            if meta.get("amount_mode") == "importe":
                amount_col = str(meta.get("amount_a") or "")
                score, partial = amount_col_score(amount_col)
                if partial:
                    problems.append(f"{side_name}: la columna elegida parece importe parcial ({amount_col}). No se procesa para evitar falsos punteos.")
                partials = [c for c, s, p, n in candidates if p]
                if partials and score < 90:
                    warnings.append(f"{side_name}: hay columnas de importe parcial detectadas ({', '.join(partials[:5])}). Revisa que el importe usado sea el total.")
                if len(candidates) > 1:
                    top_names = ', '.join(f"{c}({round(s,1)})" for c, s, p, n in candidates[:5])
                    warnings.append(f"{side_name}: columnas candidatas de importe: {top_names}")
            elif meta.get("amount_mode") == "split":
                ingreso_col = str(meta.get("amount_a") or "")
                pago_col = str(meta.get("amount_b") or "")
                if ingreso_col and is_partial_amount_col(ingreso_col):
                    problems.append(f"{side_name}: columna de ingresos sospechosa/parcial ({ingreso_col}).")
                if pago_col and is_partial_amount_col(pago_col):
                    problems.append(f"{side_name}: columna de pagos sospechosa/parcial ({pago_col}).")
                inc_cands = split_col_candidates(raw_df, SPLIT_INGRESO_NAMES)
                pay_cands = split_col_candidates(raw_df, SPLIT_PAGO_NAMES)
                inc_partials = [c for c,s,p,n,name in inc_cands if p]
                pay_partials = [c for c,s,p,n,name in pay_cands if p]
                if inc_partials or pay_partials:
                    warnings.append(f"{side_name}: columnas parciales en ingresos/pagos detectadas: {', '.join((inc_partials+pay_partials)[:6])}")
        except Exception:
            pass
        if len(df) > 0 and (meta.get("ayto_col") or meta.get("ordinal_col")):
            missing_scope = 0
            try:
                for _, r in df.head(200).iterrows():
                    if not r.get("ayuntamiento") or not r.get("ordinal"):
                        missing_scope += 1
                if missing_scope:
                    warnings.append(f"{side_name}: {missing_scope} movimientos tienen Ayuntamiento/Ordinal incompleto; esos cruces se mandarán a revisión si afectan al punteo.")
            except Exception:
                pass

    return problems, warnings, meta_left, meta_right, len(left), len(right)



def is_probably_locked(path):
    """
    ULTRA 10:
    Solo bloquea si ni siquiera se puede abrir en lectura.
    No usamos r+b como bloqueo duro porque puede dar falsos positivos en carpetas
    protegidas o archivos descargados/sin permiso de escritura.
    """
    try:
        with open(path, "rb"):
            pass
        return False
    except PermissionError:
        return True
    except Exception:
        return False

def ensure_safe_output_dir(out_dir):
    """
    Crea una carpeta de salida si no existe.
    """
    if not out_dir:
        out_dir = os.path.join(os.path.expanduser("~"), "Desktop", "Conciliador_Salidas")
    os.makedirs(out_dir, exist_ok=True)
    return out_dir

def make_backup_copy(input_path, out_dir):
    """
    Copia de seguridad del archivo original en la carpeta de salida.
    No se toca el original nunca.
    """
    backup_dir = os.path.join(out_dir, "COPIAS_SEGURIDAD_ORIGINALES")
    os.makedirs(backup_dir, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    name = os.path.basename(input_path)
    dest = os.path.join(backup_dir, f"{stamp}_{name}")
    try:
        shutil.copy2(input_path, dest)
        return dest
    except Exception:
        return ""

def write_diagnostic_report(path, data):
    """
    Informe txt legible para saber qué detectó la app.
    """
    try:
        with open(path, "w", encoding="utf-8") as f:
            f.write("DIAGNOSTICO CONCILIADOR BANCARIO ULTRA 10\n")
            f.write("=" * 55 + "\n\n")
            for k, v in data.items():
                f.write(f"{k}: {v}\n")
    except Exception:
        pass

def validate_files_safety(file_left, file_right, out_dir):
    problems = []
    warnings = []

    for label, path in [("Excel 1", file_left), ("Excel 2", file_right)]:
        if not os.path.exists(path):
            problems.append(f"{label}: el archivo no existe.")
            continue
        if os.path.isdir(path):
            problems.append(f"{label}: has seleccionado una carpeta, no un Excel.")
            continue
        ext = os.path.splitext(path)[1].lower()
        if ext not in [".xlsx", ".xlsm", ".xls"]:
            problems.append(f"{label}: formato no soportado ({ext}). Usa .xlsx, .xlsm o .xls.")
        if is_probably_locked(path):
            problems.append(f"{label}: parece estar abierto o bloqueado. Cierra Excel y vuelve a probar.")
        if os.path.basename(path).startswith("~$"):
            problems.append(f"{label}: parece ser un archivo temporal de Excel (~$). Elige el archivo real.")
        if os.path.basename(path).upper().endswith("_PUNTEADO.XLSX"):
            warnings.append(f"{label}: estás usando un archivo ya punteado. Recomendado usar el original completo.")

    try:
        out_dir = ensure_safe_output_dir(out_dir)
        test_file = os.path.join(out_dir, "_test_permiso_escritura.tmp")
        with open(test_file, "w", encoding="utf-8") as f:
            f.write("test")
        os.remove(test_file)
    except Exception:
        problems.append("Carpeta de salida: no tengo permiso de escritura. Elige otra carpeta.")

    return problems, warnings

def detect_suspicious_amount_column(meta):
    col = str(meta.get("amount_a") or meta.get("amount_b") or "")
    c = clean_text(col)
    if any(x in c for x in ["CUENTA", "ORDINAL", "NUMERO", "NIF", "DNI", "CIF", "FECHA"]):
        return True, col
    return False, col


def parse_min_common_words_value(value):
    try:
        v = int(str(value).strip() or "2")
    except Exception:
        raise ValueError("Palabras mínimas debe ser un número entero, por ejemplo 2.")
    if v < 1 or v > 6:
        raise ValueError("Palabras mínimas debe estar entre 1 y 6. Recomendado: 2.")
    return v


def secure_signature(secure):
    """
    Firma estable de los punteos para comprobar que dos ejecuciones dan lo mismo.
    """
    sig = []
    for s in secure:
        li = tuple(sorted(int(x) for x in s.get("left_indices", [])))
        ri = tuple(sorted(int(x) for x in s.get("right_indices", [])))
        sig.append((s.get("tipo",""), li, ri, round(float(s.get("importe_excel1",0)),2), round(float(s.get("importe_excel2",0)),2)))
    return sorted(sig)

def self_audit_secure_matches(secure, left, right):
    """
    Auditoría posterior. Si algo no cuadra, la app debe avisar.
    """
    problems = []
    used_left = {}
    used_right = {}

    for n, s in enumerate(secure, start=1):
        lis = list(s.get("left_indices", []))
        ris = list(s.get("right_indices", []))

        for li in lis:
            used_left.setdefault(int(li), []).append(n)
        for ri in ris:
            used_right.setdefault(int(ri), []).append(n)

        try:
            total_l = round(float(left.loc[lis, "importe"].sum()), 2)
            total_r = round(float(right.loc[ris, "importe"].sum()), 2)
            if total_l != total_r:
                problems.append(f"Punteo {n}: diferencia interna {total_l} vs {total_r}")
        except Exception as e:
            problems.append(f"Punteo {n}: no se pudo calcular suma interna ({e})")

        # En punteos uno-a-uno debe haber fecha exacta y ámbito compatible.
        if len(lis) == 1 and len(ris) == 1:
            li, ri = lis[0], ris[0]
            try:
                if not amount_equal(left.loc[li, "importe"], right.loc[ri, "importe"]):
                    problems.append(f"Punteo {n}: importe no exacto")
                if not same_day(left.loc[li, "fecha"], right.loc[ri, "fecha"]):
                    problems.append(f"Punteo {n}: fecha no exacta")
                if not same_scope(left.loc[li], right.loc[ri]):
                    problems.append(f"Punteo {n}: ámbito no compatible")
            except Exception as e:
                problems.append(f"Punteo {n}: error revisando uno-a-uno ({e})")

        # En remesas: cada línea debe tener ámbito compatible contra el otro lado si es agrupación.
        if len(lis) > 1 and len(ris) == 1:
            ri = ris[0]
            for li in lis:
                try:
                    if not same_scope(left.loc[li], right.loc[ri]):
                        problems.append(f"Punteo {n}: remesa con ámbito incompatible en Excel 1 fila índice {li}")
                except Exception:
                    pass
        if len(ris) > 1 and len(lis) == 1:
            li = lis[0]
            for ri in ris:
                try:
                    if not same_scope(left.loc[li], right.loc[ri]):
                        problems.append(f"Punteo {n}: remesa con ámbito incompatible en Excel 2 fila índice {ri}")
                except Exception:
                    pass

    duplicated_left = {k:v for k,v in used_left.items() if len(v) > 1}
    duplicated_right = {k:v for k,v in used_right.items() if len(v) > 1}
    if duplicated_left:
        problems.append(f"Filas Excel 1 usadas en más de un punteo: {duplicated_left}")
    if duplicated_right:
        problems.append(f"Filas Excel 2 usadas en más de un punteo: {duplicated_right}")

    return problems





def independent_amount_date_scope_audit(secure, left, right):
    """
    ULTRA 10:
    Auditoría independiente de sumas, duplicados, fechas y ámbito.
    Mejora: el aviso de ámbito incompleto se calcula por bloques, no por cada combinación repetida.
    """
    problems = []
    warnings = []
    used_left = set()
    used_right = set()

    for n, s in enumerate(secure, start=1):
        lis = list(s.get("left_indices", []))
        ris = list(s.get("right_indices", []))

        for li in lis:
            if int(li) in used_left:
                problems.append(f"Punteo {n}: fila de Banco usada más de una vez.")
            used_left.add(int(li))

        for ri in ris:
            if int(ri) in used_right:
                problems.append(f"Punteo {n}: fila de SicalWin usada más de una vez.")
            used_right.add(int(ri))

        try:
            total_l = round(float(left.loc[lis, "importe"].sum()), 2)
            total_r = round(float(right.loc[ris, "importe"].sum()), 2)
            if total_l != total_r:
                problems.append(f"Punteo {n}: suma no cuadra ({total_l} vs {total_r}).")
        except Exception as e:
            problems.append(f"Punteo {n}: error calculando sumas ({e}).")
            continue

        try:
            left_dates = {left.loc[i, "fecha"] for i in lis}
            right_dates = {right.loc[i, "fecha"] for i in ris}

            if len(lis) == 1 and len(ris) == 1:
                if left.loc[lis[0], "fecha"] != right.loc[ris[0], "fecha"]:
                    problems.append(f"Punteo {n}: fecha distinta en coincidencia uno-a-uno.")
            else:
                if left_dates and right_dates and not (left_dates & right_dates):
                    problems.append(f"Punteo {n}: remesa sin ninguna fecha común.")
        except Exception as e:
            problems.append(f"Punteo {n}: error auditando fechas ({e}).")

        try:
            left_block = left.loc[lis]
            right_block = right.loc[ris]
            for warn in audit_scope_completeness(left_block, right_block):
                warnings.append(f"Punteo {n}: {warn}")

            # Comprobar desigualdades reales solo cuando ambos lados informan datos.
            left_aytos = {str(left.loc[i, "ayuntamiento"]).strip() for i in lis if str(left.loc[i, "ayuntamiento"]).strip()}
            right_aytos = {str(right.loc[i, "ayuntamiento"]).strip() for i in ris if str(right.loc[i, "ayuntamiento"]).strip()}
            left_ords = {str(left.loc[i, "ordinal"]).strip() for i in lis if str(left.loc[i, "ordinal"]).strip()}
            right_ords = {str(right.loc[i, "ordinal"]).strip() for i in ris if str(right.loc[i, "ordinal"]).strip()}

            if left_aytos and right_aytos and left_aytos.isdisjoint(right_aytos):
                problems.append(f"Punteo {n}: Ayuntamiento distinto ({left_aytos} vs {right_aytos}).")
            if left_ords and right_ords and left_ords.isdisjoint(right_ords):
                problems.append(f"Punteo {n}: Ordinal distinto ({left_ords} vs {right_ords}).")
        except Exception as e:
            problems.append(f"Punteo {n}: error auditando ámbito ({e}).")

    return problems, warnings


def independent_candidate_pressure_audit(secure, left, right):
    """
    ULTRA 10:
    Auditoría de presión de candidatos bidireccional, indexada y con puntuación.
    - Comprueba alternativas Banco -> SicalWin.
    - Comprueba alternativas SicalWin -> Banco.
    - Usa fecha+importe para rendimiento.
    - Si hay alternativas, solo acepta el punteo si tiene justificación fuerte.
    """
    problems = []
    warnings = []

    right_index = build_index_by_date_amount(right)
    left_index = build_index_by_date_amount(left)

    for n, s in enumerate(secure, start=1):
        lis = list(s.get("left_indices", []))
        ris = list(s.get("right_indices", []))

        if len(lis) != 1 or len(ris) != 1:
            continue

        li, ri = lis[0], ris[0]
        lr, rr = left.loc[li], right.loc[ri]

        key_right = (lr.get("fecha"), round(float(lr.get("importe")), 2))
        key_left = (rr.get("fecha"), round(float(rr.get("importe")), 2))

        alt_right = []
        for rj in right_index.get(key_right, []):
            if rj != ri and same_scope(lr, right.loc[rj]):
                alt_right.append(rj)

        alt_left = []
        for lj in left_index.get(key_left, []):
            if lj != li and same_scope(left.loc[lj], rr):
                alt_left.append(lj)

        if alt_right or alt_left:
            score, reasons = audit_reason_score(s, lr, rr, min_common=2)

            if score >= 4:
                warnings.append(
                    f"Punteo {n}: tiene alternativas ({len(alt_left)} Banco / {len(alt_right)} SicalWin), "
                    f"pero se mantiene por justificación fuerte: {', '.join(reasons)}."
                )
            else:
                problems.append(
                    f"Punteo {n}: presión de candidatos sin justificación suficiente "
                    f"({len(alt_left)} alternativa(s) en Banco, {len(alt_right)} en SicalWin). "
                    f"Puntuación={score}; razones={', '.join(reasons) if reasons else 'ninguna'}."
                )

    return problems, warnings


def audit_real_matches(secure, left, right):
    """
    ULTRA 10: auditoría real basada en comprobaciones independientes.
    """
    problems = []
    warnings = []
    p1, w1 = independent_amount_date_scope_audit(secure, left, right)
    p2, w2 = independent_candidate_pressure_audit(secure, left, right)
    problems.extend(p1)
    problems.extend(p2)
    warnings.extend(w1)
    warnings.extend(w2)
    return problems, warnings


def _extract_problem_punteo_numbers(audit_problems):
    nums = set()
    for p in audit_problems:
        m = re.search(r"Punteo\s+(\d+)", str(p))
        if m:
            try:
                nums.add(int(m.group(1)))
            except Exception:
                pass
    return nums


def audit_scope_completeness(left_rows, right_rows):
    """
    ULTRA 10:
    Determina si el ámbito está realmente incompleto.
    No considera problema que ambos lados carezcan de Ayuntamiento/Ordinal.
    """
    left_has_ayto = any(str(r.get("ayuntamiento","")).strip() for _, r in left_rows.iterrows())
    right_has_ayto = any(str(r.get("ayuntamiento","")).strip() for _, r in right_rows.iterrows())
    left_has_ord = any(str(r.get("ordinal","")).strip() for _, r in left_rows.iterrows())
    right_has_ord = any(str(r.get("ordinal","")).strip() for _, r in right_rows.iterrows())

    warnings = []
    if left_has_ayto != right_has_ayto:
        warnings.append("Ayuntamiento informado solo en un lado.")
    if left_has_ord != right_has_ord:
        warnings.append("Ordinal/Cuenta informado solo en un lado.")
    return warnings

def build_index_by_date_amount(df):
    """
    Índice rápido por fecha+importe.
    """
    idx = {}
    for i, row in df.iterrows():
        try:
            key = (row.get("fecha"), round(float(row.get("importe")), 2))
            idx.setdefault(key, []).append(i)
        except Exception:
            continue
    return idx

def build_index_by_ref(df):
    """
    Índice de referencias reales para auditoría y explicación.
    """
    idx = {}
    for i, row in df.iterrows():
        refs = extract_refs(str(row.get("concepto", "")))
        for ref in refs:
            if is_real_reference_word(ref):
                idx.setdefault(ref, []).append(i)
    return idx

def audit_reason_score(s, lr, rr, min_common=2):
    """
    Puntuación explicativa para saber por qué un punteo se considera defendible.
    """
    points = 0
    reasons = []

    try:
        if same_tercero(lr.get("tercero",""), rr.get("tercero","")):
            points += 3
            reasons.append("tercero directo")
        elif same_tercero(lr.get("tercero",""), rr.get("concepto","")) or same_tercero(rr.get("tercero",""), lr.get("concepto","")):
            points += 2
            reasons.append("tercero en concepto")
    except Exception:
        pass

    try:
        ok, common, refs = concept_match_strong(lr.get("concepto_limpio",""), rr.get("concepto_limpio",""), min_common=min_common)
        if ok:
            points += 2
            reasons.append("concepto fuerte")
    except Exception:
        pass

    try:
        lr_refs = set(extract_refs(str(lr.get("concepto", ""))))
        rr_refs = set(extract_refs(str(rr.get("concepto", ""))))
        shared_real = sorted([x for x in (lr_refs & rr_refs) if is_real_reference_word(x)])
        if shared_real:
            points += 4
            reasons.append("referencia real compartida: " + ", ".join(shared_real[:3]))
    except Exception:
        pass

    reason_text = clean_text(str(s.get("motivo","")) + " " + str(s.get("tipo","")))
    if "TERCERO" in reason_text:
        points += 1
        reasons.append("motivo por tercero")
    if "CONCEPTO" in reason_text:
        points += 1
        reasons.append("motivo por concepto")

    return points, reasons

def process_autoverified(file_left,file_right,out_dir,name_left,name_right,allow_relation_groups,allow_unique_date_amount,allow_combo_groups,min_common_words,sheet_left=0,sheet_right=0):
    """
    ULTRA 10 - Auditoría bidireccional:
    - Ejecuta motor principal.
    - Audita los punteos de forma independiente.
    - Si hay problemas en punteos concretos, esos punteos se eliminan de "seguros"
      y pasan al informe como duda crítica.
    - No bloquea todo el proceso por un único caso problemático.
    """
    out_dir = ensure_safe_output_dir(out_dir)
    os.makedirs(out_dir, exist_ok=True)

    backup_left = make_backup_copy(file_left, out_dir)
    backup_right = make_backup_copy(file_right, out_dir)

    left, meta_left = normalize(file_left,name_left,sheet_left)
    right, meta_right = normalize(file_right,name_right,sheet_right)

    secure, review, left2, right2 = reconcile(left, right, allow_relation_groups, allow_unique_date_amount, allow_combo_groups, min_common_words)

    audit_problems, audit_warnings = audit_real_matches(secure, left2, right2)
    problem_nums = _extract_problem_punteo_numbers(audit_problems)

    final_secure = []
    removed_secure = []

    for idx, s in enumerate(secure, start=1):
        if idx in problem_nums:
            removed_secure.append((idx, s))
        else:
            final_secure.append(s)

    # Pasar punteos rechazados al informe de dudas/pendientes.
    for idx, s in removed_secure:
        left_rows = ", ".join(str(int(left2.loc[i, "excel_row"])) for i in s.get("left_indices", []))
        right_rows = ", ".join(str(int(right2.loc[i, "excel_row"])) for i in s.get("right_indices", []))
        related = [p for p in audit_problems if f"Punteo {idx}:" in str(p)]
        motivo = " | ".join(str(x) for x in related) if related else "Rechazado por auditoría bidireccional."
        add_review(
            review,
            "DUDA CRITICA - auditoría bidireccional",
            left_rows,
            right_rows,
            motivo,
            "Revisar manualmente antes de puntear."
        )

    diagnostic = os.path.join(out_dir,"DIAGNOSTICO_CONCILIADOR_ULTRA10.txt")
    write_diagnostic_report(diagnostic, {
        "Excel 1": file_left,
        "Excel 2": file_right,
        "Copia seguridad Excel 1": backup_left,
        "Copia seguridad Excel 2": backup_right,
        "Movimientos Excel 1": len(left2),
        "Movimientos Excel 2": len(right2),
        "Punteados seguros propuestos": len(secure),
        "Punteados seguros finales": len(final_secure),
        "Punteos rechazados por auditoría": len(removed_secure),
        "Registros informe": len(review),
        "Columnas Excel 1": meta_left,
        "Columnas Excel 2": meta_right,
        "Auditoría bidireccional": "OK" if not audit_problems else "CON INCIDENCIAS PARCIALES",
        "Advertencias auditoría": audit_warnings,
        "Problemas auditoría": audit_problems,
        "Reglas": "ULTRA 10: auditoría de presión bidireccional, ámbito incompleto depurado y salida parcial segura",
    })

    left_marks = []
    right_marks = []
    for s in final_secure:
        for i in s["left_indices"]:
            left_marks.append((int(left2.loc[i,"excel_row"]), int(left2.loc[i,"amount_col"])))
        for i in s["right_indices"]:
            right_marks.append((int(right2.loc[i,"excel_row"]), int(right2.loc[i,"amount_col"])))

    left_out=safe_output_path(out_dir,file_left)
    right_out=safe_output_path(out_dir,file_right)
    report=safe_named_output_path(out_dir,"INFORME_DUDAS_Y_PENDIENTES.xlsx")

    mark_only_amounts(file_left,left_out,left_marks)
    mark_only_amounts(file_right,right_out,right_marks)
    make_report(report,final_secure,review,left2,right2,name_left,name_right,meta_left,meta_right,min_common_words)

    # Informe adicional si hubo incidencias parciales
    if audit_problems:
        partial_report = safe_named_output_path(out_dir, "INFORME_AUDITORIA_PARCIAL_ULTRA10.txt")
        with open(partial_report, "w", encoding="utf-8") as f:
            f.write("INFORME DE AUDITORIA PARCIAL - ULTRA 10\n")
            f.write("="*55 + "\n\n")
            f.write("La app ha generado salidas, pero algunos punteos propuestos se han eliminado y enviado a dudas.\n\n")
            f.write("Problemas detectados:\n")
            for p in audit_problems:
                f.write("- " + str(p) + "\n")
            if audit_warnings:
                f.write("\nAdvertencias:\n")
                for w in audit_warnings:
                    f.write("- " + str(w) + "\n")

    return {
        "left_out":left_out,
        "right_out":right_out,
        "report":report,
        "diagnostic":diagnostic,
        "left_count":len(left2),
        "right_count":len(right2),
        "secure":len(final_secure),
        "secure_proposed":len(secure),
        "secure_removed":len(removed_secure),
        "review":len(review),
        "meta_left":meta_left,
        "meta_right":meta_right,
        "audit_problems":audit_problems,
        "audit_warnings":audit_warnings
    }

MANUAL_TEXT = """
MANUAL ULTRA 10 SEGURIDAD REFORZADA

Recomendado:
- Excel 1: Banco completo
- Excel 2: SicalWin completo
- Remesas activadas
- Importe+fecha únicos activado si quieres imitar tu punteo manual
- Combinaciones desactivadas salvo necesidad

La app puntea:
- Remesas por Rel. Transferencia con suma exacta
- Importe + fecha + tercero/concepto fuerte
- Opcionalmente importe + fecha con candidato único

Corrección ULTRA 10:
- Protege Ingresos/Pagos frente a columnas parciales.
- Tercero más estricto.
- Ámbito incompleto ya no se trata como compatible sin más.
- Si tercero/concepto es fuerte, no castiga duplicados sin relación.

No puntea:
- Varios candidatos
- Diferencias de céntimos
- Fechas distintas
"""

class App:
    def __init__(self, root):
        self.root=root; root.title(APP_NAME); root.geometry("1020x780")
        self.mode=tk.StringVar(value="Banco ↔ SicalWin")
        self.file_left=tk.StringVar(); self.file_right=tk.StringVar(); self.out_dir=tk.StringVar()
        self.sheet_left=tk.StringVar(value='0')
        self.sheet_right=tk.StringVar(value='0')
        self.name_left=tk.StringVar(value="Banco"); self.name_right=tk.StringVar(value="SicalWin")
        self.allow_relation_groups=tk.BooleanVar(value=True)
        self.allow_unique_date_amount=tk.BooleanVar(value=True)
        self.allow_combo_groups=tk.BooleanVar(value=False)
        self.min_common_words=tk.StringVar(value="2")
        self.build()

    def build(self):
        main=ttk.Frame(self.root,padding=16); main.pack(fill="both",expand=True)
        ttk.Label(main,text=APP_NAME,font=("Arial",18,"bold")).pack(pady=(0,4))
        ttk.Label(main,text="Versión ULTRA 10: IA reforzada con más seguridad en importes, terceros y ámbitos.").pack(pady=(0,12))
        mode_frame=ttk.LabelFrame(main,text="Modo",padding=10); mode_frame.pack(fill="x",pady=6)
        ttk.Radiobutton(mode_frame,text="Banco ↔ SicalWin",variable=self.mode,value="Banco ↔ SicalWin",command=self.update_names).pack(side="left",padx=12)
        ttk.Radiobutton(mode_frame,text="Dos Excel cualquiera",variable=self.mode,value="Dos Excel",command=self.update_names).pack(side="left",padx=12)
        files=ttk.LabelFrame(main,text="Archivos",padding=10); files.pack(fill="x",pady=6)
        self.row(files,"Excel 1:",self.file_left,self.choose_left)
        self.sheet_row(files,"Hoja Excel 1:",self.sheet_left)
        self.row(files,"Excel 2:",self.file_right,self.choose_right)
        self.sheet_row(files,"Hoja Excel 2:",self.sheet_right)
        self.row(files,"Carpeta salida:",self.out_dir,self.choose_out)
        names=ttk.LabelFrame(main,text="Nombres para el informe",padding=10); names.pack(fill="x",pady=6)
        nf=ttk.Frame(names); nf.pack(fill="x")
        ttk.Label(nf,text="Nombre Excel 1:",width=18).pack(side="left")
        ttk.Entry(nf,textvariable=self.name_left,width=25).pack(side="left",padx=5)
        ttk.Label(nf,text="Nombre Excel 2:",width=18).pack(side="left",padx=(25,0))
        ttk.Entry(nf,textvariable=self.name_right,width=25).pack(side="left",padx=5)
        rules=ttk.LabelFrame(main,text="Reglas",padding=10); rules.pack(fill="x",pady=6)
        line=ttk.Frame(rules); line.pack(fill="x")
        ttk.Checkbutton(line,text="Puntear remesas por Rel. Transferencia",variable=self.allow_relation_groups).pack(side="left",padx=(0,20))
        ttk.Checkbutton(line,text="Puntear importe+fecha únicos",variable=self.allow_unique_date_amount).pack(side="left",padx=(0,20))
        ttk.Checkbutton(line,text="Buscar agrupaciones por combinación (lento)",variable=self.allow_combo_groups).pack(side="left",padx=(0,20))
        ttk.Label(line,text="Palabras mínimas:").pack(side="left")
        ttk.Entry(line,textvariable=self.min_common_words,width=5).pack(side="left",padx=6)
        tabs=ttk.Notebook(main); tabs.pack(fill="both",expand=True,pady=10)
        tab_run=ttk.Frame(tabs,padding=8); tab_manual=ttk.Frame(tabs,padding=8)
        tabs.add(tab_run,text="Ejecutar"); tabs.add(tab_manual,text="Manual")
        tk.Button(tab_run,text="COMPARAR Y MARCAR IMPORTES",command=self.run,bg="#FFFF99",height=2,width=36).pack(pady=10)
        self.progress=ttk.Progressbar(tab_run, orient="horizontal", mode="determinate", maximum=100)
        self.progress.pack(fill="x", pady=(0,10))
        self.status=tk.Text(tab_run,height=18,wrap="word"); self.status.pack(fill="both",expand=True)
        self.log("Listo. Selecciona banco completo y SicalWin completo.")
        manual_box=tk.Text(tab_manual,wrap="word"); manual_box.pack(fill="both",expand=True)
        manual_box.insert("end",MANUAL_TEXT); manual_box.configure(state="disabled")

    def row(self,parent,label,var,command):
        frame=ttk.Frame(parent); frame.pack(fill="x",pady=4)
        ttk.Label(frame,text=label,width=16).pack(side="left")
        ttk.Entry(frame,textvariable=var).pack(side="left",fill="x",expand=True)
        ttk.Button(frame,text="Elegir",command=command).pack(side="left",padx=6)


    def sheet_row(self,parent,label,var):
        frame=ttk.Frame(parent); frame.pack(fill="x",pady=4)
        ttk.Label(frame,text=label,width=16).pack(side="left")
        ttk.Entry(frame,textvariable=var,width=30).pack(side="left")
        ttk.Label(frame,text="Nombre de hoja o 0 para la primera").pack(side="left",padx=8)

    def update_names(self):
        if self.mode.get()=="Banco ↔ SicalWin":
            self.name_left.set("Banco"); self.name_right.set("SicalWin")
        else:
            self.name_left.set("Excel1"); self.name_right.set("Excel2")

    def choose_left(self):
        p=filedialog.askopenfilename(filetypes=[("Excel","*.xlsx *.xlsm *.xls")])
        if p:
            self.file_left.set(p)
            self.warn_if_punteado(p)
            sheets=list_sheets(p)
            if sheets:
                self.sheet_left.set(sheets[0])
                self.log("Hojas Excel 1: " + ", ".join(sheets))
    def choose_right(self):
        p=filedialog.askopenfilename(filetypes=[("Excel","*.xlsx *.xlsm *.xls")])
        if p:
            self.file_right.set(p)
            self.warn_if_punteado(p)
            sheets=list_sheets(p)
            if sheets:
                self.sheet_right.set(sheets[0])
                self.log("Hojas Excel 2: " + ", ".join(sheets))
    def warn_if_punteado(self,p):
        if os.path.basename(p).upper().endswith("_PUNTEADO.XLSX"):
            messagebox.showwarning("Archivo ya punteado","Has seleccionado un archivo que ya acaba en _PUNTEADO.xlsx.\n\nLo recomendable es usar siempre el banco/SicalWin original completo.")
    def sheet_value(self,var):
        v=var.get().strip()
        if v=="" or v=="0":
            return 0
        return v
    def choose_out(self):
        p=filedialog.askdirectory()
        if p: self.out_dir.set(p)
    def log(self,msg):
        self.status.insert("end",msg+"\n"); self.status.see("end"); self.root.update_idletasks()

    def run(self):
        if not self.file_left.get() or not self.file_right.get() or not self.out_dir.get():
            messagebox.showwarning("Faltan datos","Elige los dos Excel y la carpeta de salida.")
            return

        self.status.delete("1.0","end")
        self.progress["value"] = 0
        self.log("Validando seguridad de archivos...")

        try:
            min_words_checked = parse_min_common_words_value(self.min_common_words.get())
        except Exception as e:
            messagebox.showerror("Valor incorrecto", str(e))
            self.log("No se procesa: " + str(e))
            return

        file_problems, file_warnings = validate_files_safety(self.file_left.get(), self.file_right.get(), self.out_dir.get())
        if file_warnings:
            self.log("Avisos de archivo:")
            for w in file_warnings:
                self.log("- " + w)
        if file_problems:
            self.log("No se procesa por seguridad:")
            for p in file_problems:
                self.log("- " + p)
            messagebox.showerror("No puedo procesar", "\n".join(file_problems))
            return

        self.log("Validando columnas antes de procesar...")

        try:
            sheet_l=self.sheet_value(self.sheet_left)
            sheet_r=self.sheet_value(self.sheet_right)
            problems,warnings,meta_l,meta_r,n_l,n_r=validate_before_process(self.file_left.get(),self.file_right.get(),sheet_l,sheet_r)
            self.log(f"Excel 1: {n_l} movimientos detectados")
            self.log(f"Excel 2: {n_r} movimientos detectados")
            self.log(f"Columnas Excel 1: Fecha={meta_l.get('fecha_col')} | Importe={meta_l.get('amount_a') or meta_l.get('amount_b')} | Tercero={meta_l.get('tercero_col')} | Concepto={meta_l.get('concepto_cols')}")
            self.log(f"Columnas Excel 2: Fecha={meta_r.get('fecha_col')} | Importe={meta_r.get('amount_a') or meta_r.get('amount_b')} | Tercero={meta_r.get('tercero_col')} | Concepto={meta_r.get('concepto_cols')}")
            if warnings:
                self.log("\nAvisos:")
                for w in warnings:
                    self.log("- " + w)
            if problems:
                self.log("\nNo se procesa por estos problemas:")
                for p in problems:
                    self.log("- " + p)
                messagebox.showerror("No puedo procesar","\n".join(problems))
                return
        except Exception as e:
            self.write_log(traceback.format_exc())
            messagebox.showerror("Error validando",str(e))
            return

        self.progress["value"] = 10
        self.log("\nProcesando... Si son muchos datos, espera a que termine.")

        def worker():
            try:
                min_words=min_words_checked
                result=process_autoverified(
                    self.file_left.get(),self.file_right.get(),self.out_dir.get(),
                    self.name_left.get().strip() or "Excel1",
                    self.name_right.get().strip() or "Excel2",
                    self.allow_relation_groups.get(),
                    self.allow_unique_date_amount.get(),
                    self.allow_combo_groups.get(),
                    min_words,
                    sheet_l,
                    sheet_r
                )
                self.root.after(0,lambda:self.finish_ok(result))
            except Exception as e:
                err=traceback.format_exc()
                self.write_log(err)
                self.root.after(0,lambda:self.finish_error(str(e)))

        threading.Thread(target=worker,daemon=True).start()

    def finish_ok(self,result):
        self.progress["value"] = 100
        msg=(f"Terminado.\n\n{self.name_left.get()}: {result['left_count']} movimientos leídos\n{self.name_right.get()}: {result['right_count']} movimientos leídos\nPunteados seguros: {result['secure']}\nRegistros en informe: {result['review']}\n\n"
             f"Columnas detectadas:\n- Fecha {self.name_left.get()}: {result['meta_left'].get('fecha_col')}\n- Tercero {self.name_left.get()}: {result['meta_left'].get('tercero_col')}\n- Importe {self.name_left.get()}: {result['meta_left'].get('amount_a') or result['meta_left'].get('amount_b')}\n- Concepto {self.name_left.get()}: {result['meta_left'].get('concepto_cols')}\n"
             f"- Fecha {self.name_right.get()}: {result['meta_right'].get('fecha_col')}\n- Tercero {self.name_right.get()}: {result['meta_right'].get('tercero_col')}\n- Importe {self.name_right.get()}: {result['meta_right'].get('amount_a') or result['meta_right'].get('amount_b')}\n- Concepto {self.name_right.get()}: {result['meta_right'].get('concepto_cols')}\n\n"
             f"Archivos generados:\n- {result['left_out']}\n- {result['right_out']}\n- {result['report']}")
        self.log("\n"+msg)
        messagebox.showinfo("Terminado",msg)

    def finish_error(self,msg):
        self.progress["value"] = 0
        self.log("\nERROR: "+msg)
        self.log("Detalle técnico guardado en error_conciliador.log")
        messagebox.showerror("Error",msg+"\n\nDetalle técnico guardado en error_conciliador.log")

    def write_log(self,text):
        try:
            out=self.out_dir.get() or os.getcwd()
            with open(os.path.join(out,"error_conciliador.log"),"w",encoding="utf-8") as f:
                f.write(text)
        except Exception:
            pass


if __name__=="__main__":
    root=tk.Tk(); App(root); root.mainloop()
